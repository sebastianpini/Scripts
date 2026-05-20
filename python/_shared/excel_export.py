"""
Shared helper for writing a list of dicts to a formatted .xlsx file.

Requires openpyxl:  pip install openpyxl
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_xlsx(rows, columns, path):
    """Write *rows* (list of dicts) to *path* as a formatted Excel table.

    Only the keys listed in *columns* are written, in that order.
    Column widths are fitted to content (capped at 60 characters).
    An Excel table with filters and frozen header row is created.
    """
    wb = Workbook()
    ws = wb.active

    ws.append(columns)

    for row in rows:
        ws.append([row.get(col, "") for col in columns])

    _format_header(ws)
    _autofit_columns(ws)
    _add_table(ws)
    ws.freeze_panes = "A2"

    wb.save(path)


def _format_header(ws):
    fill = PatternFill("solid", start_color="2E75B6")
    font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 20

    data_font = Font(name="Arial", size=10)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font


def _autofit_columns(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=8)
        ws.column_dimensions[letter].width = min(max_len + 3, 60)


def _add_table(ws):
    last_col = get_column_letter(ws.max_column)
    ref = f"A1:{last_col}{ws.max_row}"

    name = "Tabelle1"
    existing = {t.name for t in ws._tables.values()} if hasattr(ws._tables, "values") else set()
    i = 1
    while name in existing:
        i += 1
        name = f"Tabelle{i}"

    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
